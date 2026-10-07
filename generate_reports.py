"""Perform observed HTTP checks and create required delivery reports."""
from __future__ import annotations

import json
import time
import uuid
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from docx import Document
from openpyxl import Workbook

BASE_URL = "http://127.0.0.1:8000"
ARTIFACTS = Path("tests-artifacts")
RESULTS: list[dict[str, object]] = []


def request(method: str, path: str, description: str, payload: dict[str, object] | None = None, token: str | None = None, expected: tuple[int, ...] = (200,)) -> dict[str, object]:
    """Make one real API request and retain its observed result."""
    body = json.dumps(payload).encode() if payload is not None else None
    accept = "text/html" if path == "/docs" else "application/json"
    headers = {"Accept": accept, "X-Request-ID": f"report-{uuid.uuid4()}"}
    if body is not None:
        headers["Content-Type"] = "application/json"
    if token:
        headers["Authorization"] = f"Bearer {token}"
    try:
        with urlopen(Request(f"{BASE_URL}{path}", data=body, headers=headers, method=method), timeout=10) as response:
            status_code = response.status
            raw = response.read().decode()
    except HTTPError as error:
        status_code = error.code
        raw = error.read().decode()
    except OSError as error:
        status_code = 0
        raw = json.dumps({"error": str(error)})
    passed = status_code in expected
    RESULTS.append({"method": method, "endpoint": path, "description": description, "status": status_code or "N/A", "result": "PASS" if passed else "FAIL", "reason": raw[:300]})
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        return {}


def write_reports() -> None:
    """Write workbook and project summary from the HTTP observations."""
    ARTIFACTS.mkdir(exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "API checks"
    sheet.append(["#", "Method", "Endpoint", "Description", "Status Code", "Pass/Fail", "Reason"])
    for index, result in enumerate(RESULTS, 1):
        sheet.append([index, result["method"], result["endpoint"], result["description"], result["status"], result["result"], result["reason"]])
    workbook.save(ARTIFACTS / "api_test_report.xlsx")
    (ARTIFACTS / "test_results.json").write_text(json.dumps(RESULTS, indent=2), encoding="utf-8")

    document = Document()
    document.add_heading("Story Learning API — Verification Summary", 0)
    document.add_paragraph(f"Generated from {len(RESULTS)} live HTTP checks.")
    document.add_heading("Observed results", level=1)
    for result in RESULTS:
        document.add_paragraph(f"{result['result']}: {result['method']} {result['endpoint']} — HTTP {result['status']}")
    document.save(ARTIFACTS / "project_report.docx")


for _ in range(30):
    health = request("GET", "/health", "Public readiness endpoint", expected=(200,))
    if health.get("success"):
        break
    RESULTS.pop()
    time.sleep(1)
else:
    RESULTS.append({"method": "GET", "endpoint": "/health", "description": "Public readiness endpoint", "status": "N/A", "result": "FAIL", "reason": "30 boot polling attempts exhausted"})

request("GET", "/docs", "Public Swagger UI", expected=(200,))
request("GET", "/openapi.json", "Public OpenAPI 3 schema", expected=(200,))
request("GET", "/api/v1/schema/", "Versioned OpenAPI schema", expected=(200,))
register = request("POST", "/api/v1/auth/register", "Public account registration", {"email": f"report-{uuid.uuid4()}@example.com", "password": "Safe-report-password-123"}, expected=(201,))
token = register.get("data", {}).get("access") if isinstance(register.get("data"), dict) else None
catalogue = request("GET", "/api/v1/stories?limit=2&offset=0", "Public active story collection with pagination", expected=(200,))
request("GET", "/api/v1/stories/search?q=Little", "Public story keyword search", expected=(200,))
items = catalogue.get("data", {}).get("items", []) if isinstance(catalogue.get("data"), dict) else []
if token and items:
    story_id = items[0]["id"]
    request("POST", f"/api/v1/stories/{story_id}/favorite", "Authenticated favorite creation", token=token, expected=(201,))
    request("POST", f"/api/v1/stories/{story_id}/complete", "Authenticated idempotent completion", token=token, expected=(200,))
    request("GET", "/api/v1/favorites?limit=20&offset=0", "Authenticated scoped favorites", token=token, expected=(200,))
    request("GET", "/api/v1/progress?limit=20&offset=0", "Authenticated scoped progress", token=token, expected=(200,))
else:
    RESULTS.append({"method": "POST", "endpoint": "/api/v1/stories/{id}/favorite", "description": "Authenticated favorite creation", "status": "N/A", "result": "FAIL", "reason": "No successful token or seeded story from observed requests"})
write_reports()
