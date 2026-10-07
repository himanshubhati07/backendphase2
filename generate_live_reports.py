"""Run live HTTP checks and produce the required delivery artifacts."""
from __future__ import annotations

import json
import sys
from datetime import UTC, datetime
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font

BASE_URL = "http://127.0.0.1:8000"
ARTIFACTS = Path("tests-artifacts")
RESULTS_PATH = ARTIFACTS / "test_results.json"
XLSX_PATH = ARTIFACTS / "api_test_report.xlsx"
DOCX_PATH = ARTIFACTS / "project_report.docx"


def request(method: str, endpoint: str, description: str, payload: dict | None = None, headers: dict | None = None) -> dict:
    body = json.dumps(payload).encode() if payload is not None else None
    request_headers = {"Accept": "application/json"}
    if body is not None:
        request_headers["Content-Type"] = "application/json"
    if headers:
        request_headers.update(headers)
    req = Request(f"{BASE_URL}{endpoint}", data=body, headers=request_headers, method=method)
    try:
        with urlopen(req, timeout=10) as response:
            status_code = response.status
            raw_body = response.read().decode()
    except HTTPError as error:
        status_code = error.code
        raw_body = error.read().decode()
    try:
        response_body = json.loads(raw_body)
    except json.JSONDecodeError:
        response_body = raw_body
    passed = 200 <= status_code < 300
    return {
        "method": method,
        "endpoint": endpoint,
        "description": description,
        "status_code": status_code,
        "pass_fail": "PASS" if passed else "FAIL",
        "reason": "Live response received." if passed else f"Unexpected live HTTP status {status_code}.",
        "response": response_body,
    }


def expected_request(method: str, endpoint: str, description: str, expected_status: int, payload: dict | None = None, headers: dict | None = None) -> dict:
    result = request(method, endpoint, description, payload, headers)
    result["pass_fail"] = "PASS" if result["status_code"] == expected_status else "FAIL"
    result["reason"] = (
        "Live response matched expected status."
        if result["pass_fail"] == "PASS"
        else f"Expected {expected_status}; live response was {result['status_code']}."
    )
    return result


def main() -> int:
    ARTIFACTS.mkdir(exist_ok=True)
    results = [
        expected_request("GET", "/health", "Unauthenticated health probe", 200),
        expected_request("GET", "/docs", "Swagger UI documentation", 200),
        expected_request("GET", "/openapi.json", "OpenAPI specification", 200),
        expected_request("GET", "/api/v1/tasks", "Protected task list rejects missing bearer token", 401),
    ]
    email = "live.report@example.com"
    registration = expected_request(
        "POST", "/api/v1/auth/register", "Register account for authenticated live checks", 201,
        {"email": email, "password": "CorrectHorse1"},
    )
    results.append(registration)
    login = expected_request(
        "POST", "/api/v1/auth/login", "Authenticate live report account", 200,
        {"email": email, "password": "CorrectHorse1"},
    )
    results.append(login)
    token = login["response"].get("data", {}).get("access_token") if isinstance(login["response"], dict) else None
    if not token:
        raise RuntimeError("Live login did not return a bearer token")
    auth = {"Authorization": f"Bearer {token}"}
    results.append(expected_request(
        "POST", "/api/v1/tasks", "Create task through authenticated API", 201,
        {"title": "Live report task", "description": "Created by live report.", "priority": "High", "due_date": "2030-01-15", "status": "Pending"}, auth,
    ))
    results.append(expected_request("GET", "/api/v1/tasks", "List authenticated tasks", 200, headers=auth))

    RESULTS_PATH.write_text(json.dumps({"generated_at": datetime.now(UTC).isoformat(), "results": results}, indent=2), encoding="utf-8")
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "API Test Report"
    headers = ["#", "Method", "Endpoint", "Description", "Status Code", "Pass/Fail", "Reason"]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
    for index, result in enumerate(results, start=1):
        worksheet.append([index, result["method"], result["endpoint"], result["description"], result["status_code"], result["pass_fail"], result["reason"]])
    for column, width in {"A": 6, "B": 10, "C": 24, "D": 46, "E": 14, "F": 12, "G": 42}.items():
        worksheet.column_dimensions[column].width = width
    workbook.save(XLSX_PATH)

    document = Document()
    document.add_heading("Task Manager API – Live Verification Summary", 0)
    document.add_paragraph(f"Generated from live HTTP checks at {datetime.now(UTC).isoformat()}.")
    document.add_paragraph(f"Checks: {len(results)}; passed: {sum(item['pass_fail'] == 'PASS' for item in results)}.")
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["Method", "Endpoint", "Status", "Result"]):
        cell.text = label
    for result in results:
        row = table.add_row().cells
        row[0].text = result["method"]
        row[1].text = result["endpoint"]
        row[2].text = str(result["status_code"])
        row[3].text = result["pass_fail"]
    document.save(DOCX_PATH)
    return 0 if all(result["pass_fail"] == "PASS" for result in results) else 1


if __name__ == "__main__":
    sys.exit(main())
